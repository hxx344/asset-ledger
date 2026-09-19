"""Read the source workbook without changing it. Keep cached history and provenance."""
import datetime
import json
import pathlib
import re
import sys
import openpyxl

source = pathlib.Path(sys.argv[1])
as_of = sys.argv[2] if len(sys.argv) > 2 else datetime.date.today().isoformat()
values = openpyxl.load_workbook(source, data_only=True).worksheets[0]
formulas = openpyxl.load_workbook(source, data_only=False).worksheets[0]
def date_of(value):
    if isinstance(value, datetime.datetime):
        return value.date().isoformat()
    parts = re.findall(r'\d+', str(value))
    if len(parts) != 3:
        raise ValueError(f'Unrecognized date: {value}')
    return datetime.date(*map(int, parts)).isoformat()

periods = []
for col in range(5, 99, 3):
    date = date_of(values.cell(1, col).value)
    rows = []
    project = ''
    for row in range(2, 41):
        project = values.cell(row, 1).value or project
        kind = values.cell(row, 2).value
        if not kind:
            continue
        q = values.cell(row, col - 2).value
        p = values.cell(row, col - 1).value
        v = values.cell(row, col).value
        rows.append({'id': f'row-{row}', 'project': project, 'kind': kind,
                     'quantity': q, 'price': p, 'value': v,
                     'cell': values.cell(row, col).coordinate,
                     'formula': formulas.cell(row, col).value})
    total = values.cell(42, col).value
    calculated = sum(r['value'] or 0 for r in rows)
    periods.append({'id': values.cell(1, col).coordinate, 'date': date, 'total': total,
                    'fx': values.cell(43, col).value, 'cny': values.cell(44, col).value,
                    'rows': rows, 'future': date > as_of, 'rowSum': calculated,
                    'difference': round(calculated - total, 8)})

baseline = max((p for p in periods if not p['future']), key=lambda p: (p['date'], periods.index(p)))
if abs(baseline['difference']) > 0.01:
    raise ValueError('The baseline total does not reconcile')
data = {'source': source.name, 'sheet': values.title, 'startedAt': as_of,
        'baselineId': baseline['id'], 'baselineDate': baseline['date'],
        'fx': baseline['fx'], 'periods': periods}
target = pathlib.Path(sys.argv[3]) if len(sys.argv) > 3 else pathlib.Path('lib/imported-ledger.json')
target.parent.mkdir(parents=True, exist_ok=True)
target.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'periods': len(periods), 'rows': len(baseline['rows']), 'baseline': baseline['date'],
                  'total': baseline['total'], 'future': [p['date'] for p in periods if p['future']]}, ensure_ascii=False))
